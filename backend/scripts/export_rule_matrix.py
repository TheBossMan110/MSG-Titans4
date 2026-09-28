"""
Export the Complaint Resolution Rule Matrix to CSV and XLSX.

    cd backend
    .venv\\Scripts\\python.exe scripts\\export_rule_matrix.py            # -> reports/
    .venv\\Scripts\\python.exe scripts\\export_rule_matrix.py --out DIR

Reads the same YAML files, in the same order, that ``src/db/seed/rules.py``
loads into the ``rules`` table:

    complaint_rules/*.yaml   routing_rules/*.yaml   escalation_rules/*.yaml

(``*.zenithra.bak`` backups do not match ``*.yaml`` and are ignored, exactly as
the loader ignores them.) Nothing here touches a database or the network, so
the export can be regenerated anywhere the repository is checked out.

Outputs
-------
``rule_matrix.csv``   one row per loaded rule, columns in the order the SRS
                      lists them, then Rule type, Precedence, Version, Active.
``rule_matrix.xlsx``  sheet "Rule Matrix"  - the same rows as the CSV
                      sheet "By Case"      - one row per case: every rule that
                                              shares a lineage (RULE-002, ESC-002,
                                              RES-002, ELG-002-*) merged with the
                                              engine's own merge semantics
                      sheet "Summary"      - counts, source files, checksum

A column a rule does not set is left blank. Nothing is inferred from another
rule on the per-rule sheet; the "By Case" sheet is the only place rules are
combined, and it says which rules each row was built from.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
import sys
from collections import Counter, OrderedDict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import yaml

BACKEND = Path(__file__).resolve().parents[1]
RULE_DIRECTORIES = ("complaint_rules", "routing_rules", "escalation_rules")  # as the loader
DEFAULT_OUT = BACKEND / "reports"

COLUMNS = [
    "Rule ID",
    "Category",
    "Subcategory",
    "Conditions",
    "Department",
    "Urgency",
    "Priority",
    "Policy",
    "Escalation",
    "Required actions",
    "Prohibited actions",
    "Follow-up",
    "Rule type",
    "Precedence",
    "Version",
    "Active",
]

TYPE_ORDER = {"CLASSIFICATION": 0, "ESCALATION": 1, "RESOLUTION": 2, "ELIGIBILITY": 3, "ROUTING": 4}
URGENCY_RANK = {"LOW": 0, "MEDIUM": 1, "HIGH": 2, "CRITICAL": 3}
NAME_PREFIXES = ("Obligations for ", "Refund for ", "Replacement for ", "Compensation for ")
OPERATORS = {
    "eq": "=", "ne": "!=", "gt": ">", "gte": ">=", "lt": "<", "lte": "<=",
    "in": "in", "not_in": "not in", "contains": "contains",
    "not_contains": "does not contain", "exists": "exists", "matches": "matches",
}
LIST_SEPARATOR_CSV = " | "


# ══════════════════════════════════════════════════════════════
# reading
# ══════════════════════════════════════════════════════════════
def rule_files() -> list[Path]:
    """The loader's file list: each directory in order, files sorted."""
    found: list[Path] = []
    for directory in RULE_DIRECTORIES:
        path = BACKEND / directory
        if path.exists():
            found.extend(sorted(path.glob("*.yaml")))
    return found


def load_specs() -> tuple[list[dict[str, Any]], list[dict[str, Any]], str, list[str]]:
    """
    Return ``(specs, files, checksum, problems)``.

    ``checksum`` is computed the way the loader computes it (sha256 over the
    raw file text, in load order), so its first eight characters match the
    suffix of the live ``ruleset_version``.
    """
    specs: list[dict[str, Any]] = []
    files: list[dict[str, Any]] = []
    problems: list[str] = []
    checksum = hashlib.sha256()
    seen: set[str] = set()

    for path in rule_files():
        raw = path.read_text(encoding="utf-8")
        checksum.update(raw.encode("utf-8"))
        document = yaml.safe_load(raw) or {}
        relative = path.relative_to(BACKEND).as_posix()
        default_precedence = int(document.get("default_precedence", 50))
        rules = document.get("rules") or []
        files.append(
            {
                "file": relative,
                "ruleset": document.get("ruleset", ""),
                "version": str(document.get("version", "")),
                "rules": len(rules),
                "sha256": hashlib.sha256(raw.encode("utf-8")).hexdigest(),
            }
        )
        for spec in rules:
            if not isinstance(spec, dict):
                problems.append(f"{relative}: a rule entry is not a mapping")
                continue
            ref = str(spec.get("rule_ref", ""))
            if not ref:
                problems.append(f"{relative}: a rule has no rule_ref")
                continue
            if ref in seen:
                problems.append(f"{ref}: duplicate rule_ref (the loader would reject it)")
                continue
            seen.add(ref)
            specs.append(
                {
                    "spec": spec,
                    "file": relative,
                    "file_version": str(document.get("version", "")),
                    "default_precedence": default_precedence,
                }
            )
    return specs, files, checksum.hexdigest(), problems


# ══════════════════════════════════════════════════════════════
# rendering
# ══════════════════════════════════════════════════════════════
def render_condition(node: Any) -> str:
    """The ``when`` tree as a readable boolean expression."""
    if not isinstance(node, dict) or not node:
        return "?"
    if "always" in node:
        return "always (catch-all)" if node["always"] else "never"
    if "signal" in node:
        text = str(node["signal"])
        if "min_weight" in node:
            text += f" (weight >= {node['min_weight']})"
        return text
    for combinator, joiner in (("all_of", " AND "), ("any_of", " OR ")):
        if combinator in node:
            parts = [_wrap(render_condition(child)) for child in node[combinator] or []]
            return joiner.join(parts)
    if "none_of" in node:
        parts = [_wrap(render_condition(child)) for child in node["none_of"] or []]
        return "NOT (" + " OR ".join(parts) + ")"
    if "not" in node:
        return "NOT " + _wrap(render_condition(node["not"]))
    subject_key = "field" if "field" in node else "fact" if "fact" in node else None
    if subject_key:
        for operator, symbol in OPERATORS.items():
            if operator in node:
                value = node[operator]
                if isinstance(value, bool):
                    value = str(value).lower()
                return f"{node[subject_key]} {symbol} {value}"
    return str(node)


def _wrap(text: str) -> str:
    return f"({text})" if (" AND " in text or " OR " in text) else text


def signals_in(node: Any) -> set[str]:
    """Every signal a condition tree names."""
    found: set[str] = set()
    if isinstance(node, dict):
        if "signal" in node:
            found.add(str(node["signal"]))
        for key in ("all_of", "any_of", "none_of"):
            for child in node.get(key) or []:
                found |= signals_in(child)
        if "not" in node:
            found |= signals_in(node["not"])
    return found


def escalation_ranks() -> dict[str, int]:
    """The escalation ladder from config/taxonomy.yaml (rank 0 = none)."""
    path = BACKEND / "config" / "taxonomy.yaml"
    if not path.exists():
        return {}
    taxonomy = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return {
        str(level["code"]).upper(): int(level.get("rank", 0))
        for level in taxonomy.get("escalation_levels") or []
        if isinstance(level, dict) and level.get("code")
    }


def describe(spec: dict[str, Any]) -> str:
    """The authored prose a rule is named after, without its type prefix."""
    eligibility = spec.get("eligibility") or {}
    stated = eligibility.get("conditions") or []
    if stated:
        return "; ".join(str(item) for item in stated)
    name = " ".join(str(spec.get("name", "")).split())
    for prefix in NAME_PREFIXES:
        if name.startswith(prefix):
            return name[len(prefix):]
    return name


def section(ref: dict[str, Any]) -> str:
    doc = str(ref.get("doc_ref", "")).strip()
    sec = str(ref.get("section_ref", "") or "").strip()
    if not doc:
        return ""
    return f"{doc}-S{sec}" if sec else doc


def policy_text(spec: dict[str, Any]) -> str:
    then = spec.get("then") or {}
    parts = [section(ref) for ref in then.get("policy_refs") or [] if isinstance(ref, dict)]
    eligibility = spec.get("eligibility") or {}
    if eligibility:
        verdict = f"Eligibility {eligibility.get('type', '?')}: {eligibility.get('outcome', '?')}"
        if eligibility.get("requires_human_approval"):
            verdict += " (human approval required)"
        cited = section(eligibility.get("policy_ref") or {})
        parts.append(f"{cited} - {verdict}" if cited else verdict)
    return LIST_SEPARATOR_CSV.join(p for p in parts if p)


def department_text(then: dict[str, Any]) -> str:
    owner = then.get("department") or ""
    support = then.get("support_department") or ""
    if owner and support:
        return f"{owner}; support: {support}"
    return owner or (f"support: {support}" if support else "")


def escalation_text(spec: dict[str, Any]) -> str:
    level = (spec.get("then") or {}).get("escalation") or ""
    if level and spec.get("mandatory_escalation"):
        return f"{level} (mandatory floor)"
    return level


def follow_up_text(then: dict[str, Any]) -> str:
    # The authored key is `follow_up_required`; `follow_up` is what
    # src/db/seed/rules.py reads. Either counts as authored intent here.
    value = then.get("follow_up_required", then.get("follow_up"))
    return "Yes" if value is True else ""


def conditions_text(spec: dict[str, Any]) -> str:
    expression = render_condition(spec.get("when"))
    scope = spec.get("scope") or {}
    if scope:
        scoped = ", ".join(f"{k}={v}" for k, v in scope.items())
        expression = f"[scope {scoped}] {expression}"
    prose = describe(spec)
    return f"{expression} -- {prose}" if prose else expression


def to_row(item: dict[str, Any]) -> dict[str, Any]:
    spec = item["spec"]
    then = spec.get("then") or {}
    return {
        "Rule ID": spec["rule_ref"],
        "Category": then.get("category") or "",
        "Subcategory": then.get("subcategory") or "",
        "Conditions": conditions_text(spec),
        "Department": department_text(then),
        "Urgency": then.get("urgency") or "",
        "Priority": then.get("priority") or "",
        "Policy": policy_text(spec),
        "Escalation": escalation_text(spec),
        "Required actions": list(then.get("required_actions") or []),
        "Prohibited actions": list(then.get("prohibited_actions") or []),
        "Follow-up": follow_up_text(then),
        "Rule type": str(spec.get("rule_type", "")).upper(),
        "Precedence": int(spec.get("precedence", item["default_precedence"])),
        "Version": item["file_version"],
        "Active": "Yes" if bool(spec.get("active", True)) else "No",
        # carried for the case view and the summary, not exported as columns
        "_catch_all": bool(spec.get("catch_all", False)),
        "_mandatory": bool(spec.get("mandatory_escalation", False)),
        "_level": str(then.get("escalation") or "").upper(),
        "_expression": render_condition(spec.get("when")),
        "_signals": signals_in(spec.get("when")),
        "_file": item["file"],
        "_spec": spec,
    }


# ══════════════════════════════════════════════════════════════
# lineage (for ordering and for the case view)
# ══════════════════════════════════════════════════════════════
_NUMBERED = re.compile(r"(?:RULE|ESC|RES|ELG)-(\d{3})(?:-[A-Z]{4})?")
_CROSS_CUTTING = re.compile(r"(?:ESC-|RES-|ELG-)?((?:SAF|LEG|REG|URG|INJ|REP)-\d{4})(?:-[A-Z]{4})?")


def lineage(row: dict[str, Any]) -> tuple[int, str]:
    """
    Which authored case a rule belongs to.

    RULE-002, ESC-002, RES-002 and ELG-002-* are one case of the source
    matrix. The cross-cutting families (SAF, LEG, REG, URG, INJ, REP) group by
    their own reference; the STANCE rules by the signal they fire on.
    """
    ref = row["Rule ID"]
    match = _NUMBERED.fullmatch(ref)
    if match:
        return 0, f"RULE-{match.group(1)}"
    match = _CROSS_CUTTING.fullmatch(ref)
    if match:
        return 1, match.group(1)
    if "STANCE" in ref:
        return 2, f"STANCE: {row['_expression']}"
    return 3, ref


def sort_key(row: dict[str, Any]) -> tuple[Any, ...]:
    group, key = lineage(row)
    return group, key, TYPE_ORDER.get(row["Rule type"], 9), row["Rule ID"]


def _top(ranked: list[dict[str, Any]], column: str) -> str:
    """The value from the highest-precedence rule that sets ``column``."""
    for member in ranked:
        if member[column]:
            return member[column]
    return ""


def _union(ranked: list[dict[str, Any]], column: str) -> list[str]:
    """Every value of a list column, first-seen order, no repeats."""
    seen: list[str] = []
    for member in ranked:
        for value in member[column]:
            if value not in seen:
                seen.append(value)
    return seen


def build_cases(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """
    One row per case, merged as ``python_validation/rule_engine.py`` merges:
    scalar fields from the highest-precedence rule that sets them, action
    lists and policy references unioned, follow-up if any rule requires it,
    and the escalation raised to the highest mandatory floor.

    This is the outcome of a case's own rules firing together. At runtime
    other rules (cross-cutting, safety) may fire alongside and change it.
    """
    ranks = escalation_ranks()
    groups: OrderedDict[tuple[int, str], list[dict[str, Any]]] = OrderedDict()
    for row in rows:
        groups.setdefault(lineage(row), []).append(row)

    cases: list[dict[str, Any]] = []
    for (_, key), members in groups.items():
        ranked = sorted(members, key=lambda r: (-r["Precedence"], r["Rule ID"]))

        def top(column: str, ranked: list[dict[str, Any]] = ranked) -> str:
            return _top(ranked, column)

        def union(column: str, ranked: list[dict[str, Any]] = ranked) -> list[str]:
            return _union(ranked, column)

        # Conditions: one expression if every member shares it, otherwise
        # say which rules use which. They differ for the baseline case of a
        # subcategory, whose classification rule fires on the topic alone.
        by_expression: OrderedDict[str, list[str]] = OrderedDict()
        for member in sorted(members, key=lambda r: (TYPE_ORDER.get(r["Rule type"], 9), r["Rule ID"])):
            by_expression.setdefault(member["_expression"], []).append(member["Rule ID"])
        classification = next((m for m in members if m["Rule type"] == "CLASSIFICATION"), members[0])
        prose = describe(classification["_spec"])
        if len(by_expression) == 1:
            condition = next(iter(by_expression))
        else:
            condition = "; ".join(f"{', '.join(refs)}: {expr}" for expr, refs in by_expression.items())
        if prose:
            condition = f"{condition} -- {prose}"

        policies: list[str] = []
        for member in ranked:
            for ref in (member["_spec"].get("then") or {}).get("policy_refs") or []:
                cited = section(ref) if isinstance(ref, dict) else ""
                if cited and cited not in policies:
                    policies.append(cited)
        verdicts = []
        for member in sorted(members, key=lambda r: r["Rule ID"]):
            eligibility = member["_spec"].get("eligibility") or {}
            if eligibility:
                verdict = f"{str(eligibility.get('type', '?')).title()}: {eligibility.get('outcome', '?')}"
                if eligibility.get("requires_human_approval"):
                    verdict += " (human approval)"
                verdicts.append(verdict)
        policy = LIST_SEPARATOR_CSV.join(policies + (["Eligibility - " + "; ".join(verdicts)] if verdicts else []))

        # Highest-precedence proposal, then raised to the highest mandatory
        # floor in the case, as _apply_escalation_floor does.
        escalation = top("Escalation")
        floors = [m for m in members if m["_mandatory"] and m["_level"]]
        if floors:
            floor = max(floors, key=lambda m: ranks.get(m["_level"], 0))
            current = escalation.split(" ")[0] if escalation else ""
            if ranks.get(current, -1) < ranks.get(floor["_level"], 0):
                escalation = floor["Escalation"]

        cases.append(
            {
                "Rule ID": "; ".join(m["Rule ID"] for m in sorted(members, key=sort_key)),
                "Category": top("Category"),
                "Subcategory": top("Subcategory"),
                "Conditions": condition,
                "Department": top("Department"),
                "Urgency": top("Urgency"),
                "Priority": top("Priority"),
                "Policy": policy,
                "Escalation": escalation,
                "Required actions": union("Required actions"),
                "Prohibited actions": union("Prohibited actions"),
                "Follow-up": "Yes" if any(m["Follow-up"] == "Yes" for m in members) else "",
                "Rule type": " + ".join(
                    sorted({m["Rule type"] for m in members}, key=lambda t: TYPE_ORDER.get(t, 9))
                ),
                "Precedence": max(m["Precedence"] for m in members),
                "Version": "; ".join(sorted({m["Version"] for m in members})),
                "Active": "Yes" if all(m["Active"] == "Yes" for m in members) else "Partly",
                "_case": key,
            }
        )
    return cases


# ══════════════════════════════════════════════════════════════
# writing
# ══════════════════════════════════════════════════════════════
def _cell(value: Any, separator: str) -> Any:
    if isinstance(value, list):
        return separator.join(value)
    return value


def write_csv(rows: list[dict[str, Any]], path: Path) -> None:
    # utf-8-sig: the byte-order mark tells Excel the file is UTF-8.
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(COLUMNS)
        for row in rows:
            writer.writerow([_cell(row[column], LIST_SEPARATOR_CSV) for column in COLUMNS])


WIDTHS = {
    "Rule ID": 16, "Category": 16, "Subcategory": 26, "Conditions": 60, "Department": 28,
    "Urgency": 10, "Priority": 8, "Policy": 34, "Escalation": 30, "Required actions": 48,
    "Prohibited actions": 48, "Follow-up": 9, "Rule type": 15, "Precedence": 10,
    "Version": 8, "Active": 7,
}


def _sheet(workbook, title: str, rows: list[dict[str, Any]], *, first_width: int | None = None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = workbook.create_sheet(title)
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append([_cell(row[column], "\n") for column in COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, column in enumerate(COLUMNS, start=1):
        width = WIDTHS.get(column, 14)
        if index == 1 and first_width:
            width = first_width
        sheet.column_dimensions[get_column_letter(index)].width = width
    wrap = Alignment(vertical="top", wrap_text=True)
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = wrap
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def write_xlsx(
    rows: list[dict[str, Any]],
    cases: list[dict[str, Any]],
    summary: list[tuple[str, Any]],
    path: Path,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Rule Matrix", rows)
    _sheet(workbook, "By Case", cases, first_width=34)

    sheet = workbook.create_sheet("Summary")
    for key, value in summary:
        sheet.append([key, value])
        if key and value in ("", None):
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    sheet.column_dimensions["A"].width = 46
    sheet.column_dimensions["B"].width = 110
    for line in sheet.iter_rows():
        for cell in line:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


# ══════════════════════════════════════════════════════════════
# summary
# ══════════════════════════════════════════════════════════════
def build_summary(
    rows: list[dict[str, Any]],
    cases: list[dict[str, Any]],
    files: list[dict[str, Any]],
    checksum: str,
    problems: list[str],
) -> list[tuple[str, Any]]:
    by_type = Counter(row["Rule type"] for row in rows)
    active = sum(1 for row in rows if row["Active"] == "Yes")
    signals: set[str] = set()
    for row in rows:
        signals |= row["_signals"]
    documents = sorted(
        {m.group(1) for row in rows for m in re.finditer(r"(DOC-\d{3})", row["Policy"])}
    )
    source_cases = [c for c in cases if c["_case"].startswith("RULE-")]

    lines: list[tuple[str, Any]] = [
        ("Complaint Resolution Rule Matrix - export", ""),
        ("Generated (UTC)", datetime.now(UTC).strftime("%Y-%m-%d %H:%M")),
        ("Generated by", "backend/scripts/export_rule_matrix.py (no database access)"),
        ("Ruleset checksum (sha256, loader order)", checksum),
        ("Live ruleset_version suffix", checksum[:8]),
        ("", ""),
        ("Counts", ""),
        ("Rules exported", len(rows)),
        ("Active", active),
        ("Inactive", len(rows) - active),
    ]
    lines += [(f"  {rule_type}", count) for rule_type, count in sorted(
        by_type.items(), key=lambda item: TYPE_ORDER.get(item[0], 9))]
    lines += [
        ("Mandatory escalation floors", sum(1 for row in rows if row["_mandatory"])),
        ("Catch-all rules", sum(1 for row in rows if row["_catch_all"])),
        ("Rules with Follow-up = Yes", sum(1 for row in rows if row["Follow-up"] == "Yes")),
        ("Cases (By Case sheet)", len(cases)),
        ("  of which authored matrix cases RULE-001..", len(source_cases)),
        ("Distinct signals referenced", len(signals)),
        ("Knowledge-base documents cited", ", ".join(documents)),
        ("", ""),
        ("Source files (loader order)", ""),
    ]
    lines += [
        (f"  {f['file']}", f"ruleset={f['ruleset']} version={f['version']} rules={f['rules']} sha256={f['sha256'][:12]}")
        for f in files
    ]
    lines += [
        ("", ""),
        ("Column notes", ""),
        ("Conditions", "Rendered 'when' tree, then '--' and the authored prose the rule is named after."),
        ("Department", "Owning department; a supporting department is shown as 'support: CODE'."),
        ("Policy", "Cited sections as DOC-nnn-Sn. For ELIGIBILITY rules also the entitlement verdict."),
        ("Escalation", "'(mandatory floor)' marks a level the engine may raise but never lower."),
        ("Follow-up", "Yes where the YAML sets follow_up_required (or follow_up) to true; blank otherwise."),
        ("Version", "The 'version' field of the YAML file the rule is authored in."),
        ("Blank cells", "The rule does not set that field. Nothing is filled in from another rule on the Rule Matrix sheet."),
        ("By Case", "Rules of one lineage merged with the engine's semantics: highest precedence wins scalars, actions unioned, floor applied."),
    ]
    if problems:
        lines += [("", ""), ("Problems found", "")] + [("  -", p) for p in problems]
    return lines


# ══════════════════════════════════════════════════════════════
def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="output directory (default: backend/reports)")
    args = parser.parse_args(argv)

    specs, files, checksum, problems = load_specs()
    rows = sorted((to_row(item) for item in specs), key=sort_key)
    cases = build_cases(rows)
    summary = build_summary(rows, cases, files, checksum, problems)

    args.out.mkdir(parents=True, exist_ok=True)
    csv_path = args.out / "rule_matrix.csv"
    xlsx_path = args.out / "rule_matrix.xlsx"
    write_csv(rows, csv_path)
    write_xlsx(rows, cases, summary, xlsx_path)

    width = max(len(key) for key, _ in summary if key)
    for key, value in summary:
        if key.startswith("Column notes"):
            break
        print(f"{key:<{width}}  {value}" if key else "")
    print(f"\nwrote {csv_path}\nwrote {xlsx_path}")
    for problem in problems:
        print(f"PROBLEM: {problem}", file=sys.stderr)
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
