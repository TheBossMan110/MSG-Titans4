"""
Labelled dataset import (SRS 1.8 #2, #3).

    "Each team must author its own complaint dataset of at least 500 entries,
     and must be able to ingest a hidden dataset supplied at evaluation time."

Reads a CSV or XLSX of complaints and loads them with their ground-truth
labels, if any. Two very different jobs share this path on purpose:

* **the authored dataset** — 500 complaints with ``expected_*`` labels, used to
  score both pipelines;
* **the hidden dataset** — complaints dropped in by an evaluator, usually with
  no labels at all, which must process without a code change.

Making both the same code path is the point. A hidden-dataset importer written
separately from the one used daily is the one that fails on the day it matters.

**Ground-truth labels are written here and read only by the benchmark.**
``expected_category_code`` and its siblings sit on the complaint row because
that is the natural place for them, but nothing in either pipeline may consult
them — a model or a rule that could see the answer would be scoring its own
homework. ``tests/test_benchmark.py`` enforces that structurally by grepping
the pipeline packages.

**Import never analyses.** Rows land as complaints and nothing else happens.
Running 500 complaints through both pipelines during an HTTP request would time
out, and an import that half-succeeded would leave the dataset in a state
nobody could reason about. Analysis is the benchmark's job, batched and
resumable.
"""

from __future__ import annotations

import csv
import io
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from complaint_processing.entities import parse_amount
from complaint_processing.preprocess import preprocess
from src.core.logging import get_logger
from src.db.enums import Channel, ComplaintStatus
from src.db.models import Complaint, Customer

log = get_logger("services.dataset")

# The columns the authoring brief asks for.
TEXT_COLUMNS = ("title", "description")
OPTIONAL_COLUMNS = (
    "order_ref", "transaction_ref", "product", "amount", "currency",
    "channel", "customer_email", "customer_name", "requested_resolution",
    "customer_type", "external_ref", "previous_ref",
)

# Rows flushed together. Each flush is a round trip to the database; one per
# row made a 500-row import take minutes against a hosted database.
FLUSH_EVERY = 50

# Ground truth. Read by the benchmark and by nothing else.
LABEL_COLUMNS = {
    "expected_category": "expected_category_code",
    "expected_category_code": "expected_category_code",
    "expected_subcategory": "expected_subcategory_code",
    "expected_subcategory_code": "expected_subcategory_code",
    "expected_department": "expected_department_code",
    "expected_department_code": "expected_department_code",
    "expected_urgency": "expected_urgency",
    "expected_priority": "expected_priority_code",
    "expected_priority_code": "expected_priority_code",
    "expected_escalation": "expected_escalation_code",
    "expected_escalation_code": "expected_escalation_code",
}


@dataclass(slots=True)
class ImportResult:
    """What one import did, and what it refused."""

    dataset_tag: str
    imported: int = 0
    skipped: int = 0
    labelled: int = 0
    errors: list[dict[str, Any]] = field(default_factory=list)
    public_refs: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.imported > 0

    def summary(self) -> dict[str, Any]:
        return {
            "dataset_tag": self.dataset_tag,
            "imported": self.imported,
            "skipped": self.skipped,
            "labelled": self.labelled,
            "unlabelled": self.imported - self.labelled,
            "errors": self.errors[:25],
            "error_count": len(self.errors),
        }


# ══════════════════════════════════════════════════════════════
# reading
# ══════════════════════════════════════════════════════════════
def read_rows(payload: bytes, filename: str) -> list[dict[str, Any]]:
    """
    Parse a CSV or XLSX into dictionaries.

    Header names are lower-cased and stripped, so ``Order Ref``, ``order_ref``
    and ``ORDER REF`` all land on the same column. An evaluator handing over a
    spreadsheet should not have to match our capitalisation.
    """
    suffix = Path(filename).suffix.lower()

    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(payload)
    return _read_csv(payload)


def _normalise_key(key: Any) -> str:
    return str(key or "").strip().lower().replace(" ", "_").replace("-", "_")


def _read_csv(payload: bytes) -> list[dict[str, Any]]:
    # utf-8-sig strips the BOM Excel writes; without it the first header
    # becomes "﻿title" and every row silently loses its title.
    text = payload.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [
        {_normalise_key(k): v for k, v in row.items() if k is not None}
        for row in reader
    ]


def _read_xlsx(payload: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheet = book.active

    rows = sheet.iter_rows(values_only=True)
    try:
        header = [_normalise_key(cell) for cell in next(rows)]
    except StopIteration:
        return []

    out: list[dict[str, Any]] = []
    for row in rows:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        out.append(dict(zip(header, row, strict=False)))
    return out


# ══════════════════════════════════════════════════════════════
# importing
# ══════════════════════════════════════════════════════════════
def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _upper(value: Any) -> str | None:
    text = _clean(value)
    return text.upper() if text else None


def _amount(value: Any) -> float | None:
    """
    Parse a money value from a spreadsheet cell.

    Delegates to the entity extractor's parser rather than stripping
    non-digits here. A local copy of this got "Rs. 42,500" wrong in exactly
    the way the shared one was fixed for: filtering to digits-and-dots leaves
    the dot from "Rs." in front, and the amount silently becomes 0.42500 --
    a 42,500-rupee dispute that skips every value-threshold escalation rule.
    """
    text = _clean(value)
    if not text:
        return None
    parsed = parse_amount(text)
    return float(parsed) if parsed is not None else None


def _customer(db: Session, email: str | None, name: str | None) -> Customer | None:
    if not email:
        return None
    normalised = email.strip().lower()
    existing = db.execute(
        select(Customer).where(func.lower(Customer.email) == normalised)
    ).scalars().first()
    if existing is not None:
        return existing

    count = db.execute(select(func.count()).select_from(Customer)).scalar_one()
    customer = Customer(
        external_ref=f"CUST-{count + 1:06d}",
        display_name=(name or "").strip() or normalised.split("@")[0],
        email=normalised,
    )
    db.add(customer)
    db.flush()
    return customer


def import_rows(
    db: Session,
    rows: list[dict[str, Any]],
    *,
    dataset_tag: str,
    replace: bool = False,
) -> ImportResult:
    """
    Load rows as complaints, with labels where present.

    Nothing is analysed. Rows become complaints and stop there — running 500
    of them through both pipelines inside a request would time out, and a
    half-finished import is a dataset nobody can reason about.

    ``replace`` clears the tag first, so re-importing a corrected file gives
    500 complaints rather than 1,000.
    """
    tag = dataset_tag.strip().upper()
    result = ImportResult(dataset_tag=tag)

    if replace:
        removed = db.query(Complaint).filter(Complaint.dataset_tag == tag).delete()
        db.flush()
        log.info("dataset_cleared", dataset_tag=tag, removed=removed)

    # Numbered from one count, not one per row. Same scheme as intake's
    # next_public_ref, so imported and submitted complaints share a sequence.
    number = db.execute(select(func.count()).select_from(Complaint)).scalar_one()
    # The file's own ids, so a row can point at an earlier row: the corpus
    # marks a follow-up with ``previous_ref`` and repeat detection reads it.
    by_external: dict[str, Complaint] = {}
    pending = 0

    for index, row in enumerate(rows, start=2):  # row 1 is the header
        title = _clean(row.get("title"))
        description = _clean(row.get("description"))

        if not description:
            result.skipped += 1
            result.errors.append(
                {"row": index, "error": "no description", "title": title}
            )
            continue

        prepared = preprocess(description)
        customer = _customer(
            db, _clean(row.get("customer_email")), _clean(row.get("customer_name"))
        )

        channel = _upper(row.get("channel")) or Channel.IMPORT
        if channel not in {member.value for member in Channel}:
            channel = Channel.IMPORT

        number += 1
        previous = by_external.get((_clean(row.get("previous_ref")) or "").upper())
        complaint = Complaint(
            id=uuid.uuid4(),
            public_ref=f"CMP-{number:06d}",
            customer_id=customer.id if customer else None,
            customer_type=(_clean(row.get("customer_type")) or None) and _clean(row.get("customer_type"))[:32],
            previous_complaint_id=previous.id if previous else None,
            title=(title or description[:80]),
            description_raw=prepared.raw,
            description_clean=prepared.clean,
            product=_clean(row.get("product")),
            order_ref=_clean(row.get("order_ref")),
            transaction_ref=_clean(row.get("transaction_ref")),
            amount=_amount(row.get("amount")),
            currency=_upper(row.get("currency")),
            channel=channel,
            requested_resolution=_clean(row.get("requested_resolution")),
            status=ComplaintStatus.NEW,
            dataset_tag=tag,
        )

        # ── ground truth, if the file carries any ──
        labelled = False
        for column, attribute in LABEL_COLUMNS.items():
            value = _upper(row.get(column))
            if value:
                setattr(complaint, attribute, value)
                labelled = True

        db.add(complaint)
        external = (_clean(row.get("external_ref")) or "").upper()
        if external:
            by_external[external] = complaint
        pending += 1
        if pending >= FLUSH_EVERY:
            db.flush()
            pending = 0

        result.imported += 1
        result.labelled += int(labelled)
        result.public_refs.append(complaint.public_ref)

    db.flush()
    log.info(
        "dataset_imported",
        dataset_tag=tag, imported=result.imported,
        labelled=result.labelled, skipped=result.skipped,
    )
    return result


def import_file(
    db: Session,
    payload: bytes,
    filename: str,
    *,
    dataset_tag: str,
    replace: bool = False,
) -> ImportResult:
    """Import a CSV or XLSX file."""
    return import_rows(
        db, read_rows(payload, filename), dataset_tag=dataset_tag, replace=replace
    )


# ══════════════════════════════════════════════════════════════
# inspection
# ══════════════════════════════════════════════════════════════
def describe(db: Session, dataset_tag: str) -> dict[str, Any]:
    """
    What is in one dataset, and how much of it is labelled.

    The unlabelled count matters: an unlabelled complaint can be *processed*
    but cannot be *scored*, and a benchmark quietly running over a mostly
    unlabelled set would report an accuracy figure drawn from a handful of
    rows.
    """
    tag = dataset_tag.strip().upper()

    total = db.execute(
        select(func.count()).select_from(Complaint).where(Complaint.dataset_tag == tag)
    ).scalar_one()

    labelled = db.execute(
        select(func.count())
        .select_from(Complaint)
        .where(
            Complaint.dataset_tag == tag,
            Complaint.expected_category_code.isnot(None),
        )
    ).scalar_one()

    analysed = db.execute(
        select(func.count())
        .select_from(Complaint)
        .where(Complaint.dataset_tag == tag, Complaint.analyzed_at.isnot(None))
    ).scalar_one()

    return {
        "dataset_tag": tag,
        "total": total,
        "labelled": labelled,
        "unlabelled": total - labelled,
        "analysed": analysed,
        "pending_analysis": total - analysed,
    }


def datasets(db: Session) -> list[dict[str, Any]]:
    """Every dataset tag present, with its counts."""
    tags = db.execute(
        select(Complaint.dataset_tag)
        .where(Complaint.dataset_tag.isnot(None))
        .group_by(Complaint.dataset_tag)
    ).scalars().all()
    return [describe(db, tag) for tag in sorted(tags)]


def clear(db: Session, dataset_tag: str) -> int:
    """Remove a dataset. Used between benchmark runs and by the tests."""
    tag = dataset_tag.strip().upper()
    removed = db.query(Complaint).filter(Complaint.dataset_tag == tag).delete()
    db.flush()
    log.warning("dataset_cleared", dataset_tag=tag, removed=removed)
    return removed


def unanalysed(
    db: Session, dataset_tag: str, *, limit: int | None = None
) -> list[uuid.UUID]:
    """
    Complaint ids in a dataset that have not been analysed yet.

    The benchmark uses this to resume: a run interrupted at complaint 300 of
    500 picks up at 301 rather than starting again, which on a free-tier API
    is the difference between finishing and not.
    """
    query = (
        select(Complaint.id)
        .where(
            Complaint.dataset_tag == dataset_tag.strip().upper(),
            Complaint.analyzed_at.is_(None),
        )
        .order_by(Complaint.created_at)
    )
    if limit:
        query = query.limit(limit)
    return list(db.execute(query).scalars())
